from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_chroma import Chroma
from dotenv import load_dotenv
from llm_utils import generate_chain
import os
import csv
import json
import hashlib
from pathlib import Path
from docx import Document as WordDocument
from pptx import Presentation
from openpyxl import load_workbook
from bs4 import BeautifulSoup
from langchain_core.documents import Document

load_dotenv()

def load_file_documents(file_path):
    path = Path(file_path)
    ext = path.suffix.lower()

    metadata = {
        "source": str(path.resolve()),
        "file_name": path.name,
        "file_type": ext.lstrip("."),
    }

    if ext == ".pdf":
        docs = PyPDFLoader(str(path)).load()
        for doc in docs:
            doc.metadata = {
                k: v for k, v in doc.metadata.items()
                if v is not None
                and isinstance(v, (str, int, float, bool))
            }
            doc.metadata.update(metadata)
        return docs

    elif ext == ".docx":
        word = WordDocument(str(path))
        parts = [
            p.text.strip() for p in word.paragraphs
            if p.text.strip()
        ]

        for table in word.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(parts)

    elif ext == ".pptx":
        presentation = Presentation(str(path))
        docs = []

        for slide_num, slide in enumerate(presentation.slides, start=1):
            parts = []

            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text.strip())

                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(cell.text for cell in row.cells)
                        )

            if parts:
                docs.append(
                    Document(
                        page_content="\n".join(parts),
                        metadata={**metadata, "slide": slide_num},
                    )
                )

        return docs

    elif ext in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig", errors="replace")

    elif ext == ".csv":
        with path.open(
            "r", encoding="utf-8-sig", newline=""
        ) as file:
            sample = file.read(8192)
            file.seek(0)

            # Detect whether the CSV has a header row.
            # Sniffer is heuristic, so this detection is not perfect.
            try:
                has_header = csv.Sniffer().has_header(sample)
            except csv.Error:
                has_header = True

            rows = [
                [cell.strip() for cell in row]
                for row in csv.reader(file)
            ]

        # Ignore completely blank rows.
        rows = [
            row for row in rows
            if any(value for value in row)
        ]

        if not rows:
            return []

        if has_header:
            raw_headers = rows[0]
            data_rows = rows[1:]
        else:
            raw_headers = []
            data_rows = rows

        # Support inconsistent row lengths too.
        column_count = max(len(row) for row in rows)

        headers = []
        header_counts = {}

        for i in range(column_count):
            label = (
                raw_headers[i].strip()
                if i < len(raw_headers) and raw_headers[i].strip()
                else f"column_{i + 1}"
            )

            header_counts[label] = header_counts.get(label, 0) + 1

            if header_counts[label] > 1:
                label = f"{label}_{header_counts[label]}"

            headers.append(label)

        docs = []

        for record_index, row in enumerate(data_rows, start=1):
            values = row + [""] * (column_count - len(row))

            fields = [
                f"{headers[i]}: {value}"
                for i, value in enumerate(values)
                if value
            ]

            if not fields:
                continue

            docs.append(
                Document(
                    page_content=(
                        f"Record index: {record_index}\n"
                        + "\n".join(fields)
                    ),
                    metadata={
                        **metadata,
                        "record_index": record_index,
                    },
                )
            )

        return docs

    elif ext == ".xlsx":
        workbook = load_workbook(
            str(path), read_only=True, data_only=True
        )
        docs = []

        for sheet in workbook.worksheets:
            rows = [
                " | ".join(
                    "" if value is None else str(value)
                    for value in row
                )
                for row in sheet.iter_rows(values_only=True)
            ]

            sheet_text = "\n".join(
                row for row in rows if row.strip(" |")
            )

            if sheet_text.strip():
                docs.append(
                    Document(
                        page_content=sheet_text,
                        metadata={**metadata, "sheet": sheet.title},
                    )
                )

        workbook.close()
        return docs

    elif ext in {".html", ".htm"}:
        soup = BeautifulSoup(
            path.read_text(encoding="utf-8", errors="replace"),
            "html.parser",
        )

        for tag in soup(["script", "style"]):
            tag.decompose()

        text = soup.get_text(separator=" ", strip=True)

    elif ext == ".json":
        with path.open("r", encoding="utf-8-sig") as file:
            data = json.load(file)

        text = json.dumps(data, ensure_ascii=False, indent=2)

    else:
        raise ValueError(f"Unsupported file type: {ext}")

    return [
        Document(page_content=text, metadata=metadata)
    ] if text.strip() else []






def create_rag(file_paths):

    if isinstance(file_paths, (str, Path)):
        file_paths = [file_paths]

    embeds = GoogleGenerativeAIEmbeddings(
        model="models/gemini-embedding-001"
    )

    vectors = Chroma(
        collection_name="my_rag_multiformat",
        persist_directory="./chroma_db",
        embedding_function=embeds,
    )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=90,
    )

    for file_path in file_paths:
        path = Path(file_path).resolve()
        source_path = str(path)

        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        docs = load_file_documents(path)

        if not docs:
            print(f"No readable text: {path.name}")
            continue

        # Create a normalized content fingerprint
        normalized_text = " ".join(
            " ".join(doc.page_content for doc in docs).split()
        ).casefold()

        if not normalized_text:
            print(f"No text to index: {path.name}")
            continue

        content_hash = hashlib.sha256(
            normalized_text.encode("utf-8")
        ).hexdigest()

        # Find chunks already associated with this exact file path
        source_records = vectors.get(
            where={"source": source_path},
            include=["metadatas"],
        )

        old_ids = source_records["ids"]
        old_metadata = source_records["metadatas"] or []

        # If this file is unchanged, reuse its existing chunks
        if old_ids and all(
            metadata.get("content_hash") == content_hash
            for metadata in old_metadata
        ):
            print(f"Already indexed: {path.name}")
            continue

        # Keep old chunks until the replacement is indexed successfully.
        if old_ids:
            print(f"File changed; preparing safe update: {path.name}")

        # Detect identical content in other files.
        duplicate_ids = vectors.get(
            where={"content_hash": content_hash},
            limit=1,
        )["ids"]

        if duplicate_ids:
            print(
                f"Duplicate content detected. "
                f"Indexing {path.name} separately."
            )

        # Add the content hash before splitting.
        for doc in docs:
            doc.metadata["content_hash"] = content_hash

        chunks = splitter.split_documents(docs)
        for chunk_index, chunk in enumerate(chunks):
            chunk.metadata["chunk_index"] = chunk_index

        if not chunks:
            print(f"No chunks created: {path.name}")
            continue

        # First, index the new content.
        vectors.add_documents(chunks)

        # Only after successful indexing, remove the old chunks.
        if old_ids:
            vectors.delete(ids=old_ids)
            print(f"Replaced old chunks: {path.name}")

        print(f"Indexed: {path.name}")

    return vectors.as_retriever(
        search_kwargs={"k": 10}
    )






def get_record_by_index(file_path, record_index):
    if isinstance(record_index, bool) or not isinstance(
        record_index, int
    ) or record_index < 1:
        raise ValueError("record_index must be a positive integer")

    source_path = str(Path(file_path).resolve())

    # Metadata lookup does not require embedding a query.
    vectors = Chroma(
        collection_name="my_rag_multiformat",
        persist_directory="./chroma_db",
    )

    results = vectors.get(
        where={
            "$and": [
                {"source": source_path},
                {"record_index": record_index},
            ]
        },
        include=["documents", "metadatas"],
    )

    if not results["ids"]:
        return None

    chunks = sorted(
        zip(results["documents"], results["metadatas"]),
        key=lambda item: item[1].get("chunk_index", 0),
    )

    return "\n".join(document for document, _ in chunks)






def delete_document(file_path):
    source_path = str(Path(file_path).resolve())

    vectors = Chroma(
        collection_name="my_rag_multiformat",
        persist_directory="./chroma_db",
    )

    records = vectors.get(
        where={"source": source_path}
    )

    ids = records["ids"]

    if not ids:
        print(f"No indexed chunks found for: {source_path}")
        return 0

    vectors.delete(ids=ids)

    print(f"Deleted {len(ids)} chunks for: {Path(source_path).name}")
    return len(ids)





def ask_question(query,retriever):
    #retrieve chunks
    retrieved_docs=retriever.invoke(query)

    #join the retieved chunks(doc object) into text object for gemini model 
    cont="\n\n".join(
        doc.page_content for doc in retrieved_docs
    )

    #create the chain(prompt->llm)
    chain = generate_chain()

    #generate answer and get the response
    response=chain.invoke({
        "context":cont,
        "question":query
    })

    return response.content


    
def inspect_sources():
    from langchain_google_genai import GoogleGenerativeAIEmbeddings
    from langchain_chroma import Chroma

    embeds = GoogleGenerativeAIEmbeddings(
        model="models/gemini-embedding-001"
    )

    vectors = Chroma(
        collection_name="my_rag_multiformat",
        persist_directory="./chroma_db",
        embedding_function=embeds,
    )

    data = vectors.get(include=["metadatas"])

    for doc_id, metadata in zip(data["ids"], data["metadatas"]):
        print(
            doc_id,
            metadata.get("file_name"),
            metadata.get("content_hash"),
            metadata.get("source"),
        )


if __name__ == "__main__":
    inspect_sources()